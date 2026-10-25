
import csv
import io
import json
from pathlib import Path

from openpyxl import load_workbook

from question_bank.serializers import QuestionSerializer


REQUIRED_HEADERS = [
    "text",
    "question_type",
    "options",
    "correct_answer",
    "subject",
    "topic",
    "difficulty",
]


def _read_csv(file_content):
    try:
        text = file_content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV file must use UTF-8 encoding.") from exc

    reader = csv.DictReader(io.StringIO(text, newline=""))

    if reader.fieldnames is None:
        raise ValueError("The file is empty or has no header row.")

    records = []

    for row in reader:
        row_number = reader.line_num

        if None in row:
            records.append({
                **row,
                "__extra_columns__": row[None],
                "__row_number__": row_number,
            })
        else:
            records.append({
                **row,
                "__row_number__": row_number,
            })

    return reader.fieldnames, records


def _read_excel(file_content):
    try:
        workbook = load_workbook(
            io.BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "The Excel file is invalid or corrupted."
        ) from exc

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            raise ValueError(
                "The file is empty or has no header row."
            )

        headers = [
            str(value).strip() if value is not None else ""
            for value in header_row
        ]

        records = []

        for row_number, row in enumerate(rows, start=2):
            values = list(row)

            # Ignore completely empty spreadsheet rows.
            if all(value is None for value in values):
                continue

            record = {
                headers[index]: values[index]
                for index in range(min(len(headers), len(values)))
                if headers[index]
            }

            record["__row_number__"] = row_number

            # Detect rows containing more values than the header.
            if len(values) > len(headers):
                record["__extra_columns__"] = values[len(headers):]

            records.append(record)

        return headers, records

    finally:
        workbook.close()


def parse_question_file(uploaded_file):
    extension = Path(uploaded_file.name).suffix.lower()
    file_content = uploaded_file.read()

    if extension == ".csv":
        headers, rows = _read_csv(file_content)
    elif extension == ".xlsx":
        headers, rows = _read_excel(file_content)
    else:
        raise ValueError(
            "Unsupported file type. Upload a .csv or .xlsx file."
        )

    normalized_headers = [
        str(header).strip() if header is not None else ""
        for header in headers
    ]

    if len(normalized_headers) != len(set(normalized_headers)):
        raise ValueError(
            "The file contains duplicate column headers."
        )

    missing_headers = [
        header
        for header in REQUIRED_HEADERS
        if header not in normalized_headers
    ]

    if missing_headers:
        raise ValueError(
            "Missing required columns: " + ", ".join(missing_headers)
        )

    validated_questions = []
    errors = []

    for fallback_row_number, row in enumerate(rows, start=2):
        row_number = row.get(
            "__row_number__",
            fallback_row_number,
        )

        # Reject malformed rows containing extra values.
        if "__extra_columns__" in row:
            errors.append({
                "row": row_number,
                "errors": {
                    "file": [
                        "This row contains more values than the header."
                    ]
                },
            })
            continue

        normalized_row = {
            str(key).strip(): value
            for key, value in row.items()
            if key is not None and not key.startswith("__")
        }

        # Ignore completely empty rows.
        if all(
            value is None or str(value).strip() == ""
            for value in normalized_row.values()
        ):
            continue

        options_value = normalized_row.get("options")

        if options_value is None or str(options_value).strip() == "":
            options = []

        elif isinstance(options_value, list):
            options = options_value

        else:
            try:
                options = json.loads(str(options_value))
            except (json.JSONDecodeError, TypeError):
                errors.append({
                    "row": row_number,
                    "errors": {
                        "options": [
                            "Must be a valid JSON array, such as "
                            '["Option A", "Option B"], or [].'
                        ]
                    },
                })
                continue

        if not isinstance(options, list):
            errors.append({
                "row": row_number,
                "errors": {
                    "options": ["Must be a JSON array."]
                },
            })
            continue

        data = {
            field: normalized_row.get(field)
            for field in REQUIRED_HEADERS
        }

        data["options"] = options

        # Convert non-string spreadsheet values to strings before
        # passing the data to the existing serializer.
        for field in (
            "text",
            "question_type",
            "correct_answer",
            "subject",
            "topic",
            "difficulty",
        ):
            value = data[field]

            if value is not None and not isinstance(value, str):
                data[field] = str(value)

        serializer = QuestionSerializer(data=data)

        if serializer.is_valid():
            validated_questions.append(serializer.validated_data)
        else:
            errors.append({
                "row": row_number,
                "errors": serializer.errors,
            })

    if not validated_questions and not errors:
        raise ValueError("The file contains no question rows.")

    return validated_questions, errors
